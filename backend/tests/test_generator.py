"""Generator: persist slide_plan, render pptx + slide_map (PRD §7.4)."""

from __future__ import annotations

import json
from pathlib import Path

import yaml
from pptx import Presentation

from app.generation.generator import generate_deck, load_slide_map
from app.models.blueprint import Blueprint
from app.models.claim import ClaimLedger
from app.models.run import Brief
from app.models.slide import SlidePlan, SlidePlanSlide
from app.paths import BLUEPRINTS_DIR, REPO_ROOT

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _load_blueprint() -> Blueprint:
    path = BLUEPRINTS_DIR / "msl_physician_8.yaml"
    payload = yaml.safe_load(path.read_text(encoding="utf-8"))
    return Blueprint.model_validate(payload)


def _load_ledger() -> ClaimLedger:
    return ClaimLedger.model_validate(json.loads((FIXTURES / "ledger.json").read_text(encoding="utf-8")))


def complete_plan_from_fixtures() -> SlidePlan:
    """Nine-role plan that only uses frozen fixture claim IDs."""
    return SlidePlan(
        blueprint_id="msl_physician_8",
        skill_version="v1",
        slides=[
            SlidePlanSlide(
                slide=1,
                role="title",
                headline="Demo Trial — HFrEF outcomes",
                claim_ids=["C-001"],
                speaker_notes="C-001 p.2",
                required_content_filled=["paper_title", "audience"],
            ),
            SlidePlanSlide(
                slide=2,
                role="unmet_need",
                headline="Clinical context in adults with HFrEF",
                claim_ids=["C-001"],
                speaker_notes="C-001 p.2",
                required_content_filled=["clinical_context"],
            ),
            SlidePlanSlide(
                slide=3,
                role="study_design",
                headline="Randomized, double-blind, placebo-controlled design",
                claim_ids=["C-001"],
                speaker_notes="C-001 p.2",
                required_content_filled=["design_type"],
            ),
            SlidePlanSlide(
                slide=4,
                role="population",
                headline="Adults with HFrEF",
                claim_ids=["C-001"],
                speaker_notes="C-001 p.2",
                required_content_filled=["inclusion"],
            ),
            SlidePlanSlide(
                slide=5,
                role="primary_endpoint",
                headline="Primary endpoint met — HR 0.78",
                claim_ids=["C-014"],
                speaker_notes="C-014 p.4",
                required_content_filled=["primary_result", "CI_or_p"],
            ),
            SlidePlanSlide(
                slide=6,
                role="key_secondary",
                headline="Additional efficacy from the primary analysis",
                claim_ids=["C-014"],
                speaker_notes="C-014 p.4",
                required_content_filled=["secondary_results_labelled"],
            ),
            SlidePlanSlide(
                slide=7,
                role="safety",
                headline="Serious AEs in 18.2% of treated patients",
                claim_ids=["C-022"],
                speaker_notes="C-022 p.7",
                required_content_filled=["ae_profile"],
            ),
            SlidePlanSlide(
                slide=8,
                role="limitations_relevance",
                headline="Limitations and clinical relevance",
                claim_ids=["C-001"],
                speaker_notes="C-001 p.2",
                required_content_filled=["limitations", "clinical_relevance"],
            ),
            SlidePlanSlide(
                slide=9,
                role="references",
                headline="References",
                claim_ids=["C-001", "C-014", "C-022"],
                speaker_notes="C-001 p.2; C-014 p.4; C-022 p.7",
                required_content_filled=["citations"],
            ),
        ],
    )


def test_blueprint_yaml_has_prd_roles() -> None:
    blueprint = _load_blueprint()
    assert blueprint.id == "msl_physician_8"
    roles = [s.role for s in blueprint.slides]
    assert roles == [
        "title",
        "unmet_need",
        "study_design",
        "population",
        "primary_endpoint",
        "key_secondary",
        "safety",
        "limitations_relevance",
        "references",
    ]
    refs = blueprint.slides[-1]
    assert refs.counts_against_slide_count is False
    design = next(s for s in blueprint.slides if s.role == "study_design")
    assert design.required_content == ["design_type", "N", "arms", "duration", "endpoints"]
    assert design.max_claims == 6
    prompt = REPO_ROOT / "prompts" / "slide_plan.md"
    assert prompt.is_file()
    text = prompt.read_text(encoding="utf-8")
    assert "python-pptx" in text
    assert "SlidePlan" in text


def test_generate_deck_writes_artifacts_and_notes(tmp_path: Path) -> None:
    result = generate_deck(
        blueprint=_load_blueprint(),
        ledger=_load_ledger(),
        brief=Brief(),
        skill_version="v1",
        slide_plan=complete_plan_from_fixtures(),
        output_dir=tmp_path,
    )
    assert Path(result.deck_path).is_file()
    assert Path(result.slide_map_path).is_file()
    assert Path(result.slide_plan_path).is_file()
    assert Path(result.layout_spec_path).is_file()
    assert result.layout_spec is not None
    assert result.layout_spec.slides[0].elements[0].x == 0.5

    loaded = load_slide_map(result.slide_map_path)
    assert loaded.entries
    assert all(e.claim_ids for e in loaded.entries if e.text.strip())

    prs = Presentation(result.deck_path)
    assert len(prs.slides) == 9
    notes = [s.notes_slide.notes_text_frame.text for s in prs.slides]
    assert any("C-014 p.4" in n for n in notes)
    assert any("C-022 p.7" in n for n in notes)
    assert any("C-001 p.2" in n for n in notes)
    assert 2 in result.speaker_notes_pages
    assert 4 in result.speaker_notes_pages
    assert 7 in result.speaker_notes_pages


def test_locked_slides_copied_unchanged(tmp_path: Path) -> None:
    first_dir = tmp_path / "r1"
    second_dir = tmp_path / "r2"
    plan = complete_plan_from_fixtures()
    first = generate_deck(
        blueprint=_load_blueprint(),
        ledger=_load_ledger(),
        brief=Brief(),
        skill_version="v1",
        slide_plan=plan,
        output_dir=first_dir,
    )
    original = Presentation(first.deck_path).slides[0].shapes[0].text_frame.text

    mutated = plan.model_copy(deep=True)
    mutated.slides[0].headline = "THIS HEADLINE MUST NOT APPEAR"
    second = generate_deck(
        blueprint=_load_blueprint(),
        ledger=_load_ledger(),
        brief=Brief(),
        skill_version="v1",
        slide_plan=mutated,
        output_dir=second_dir,
        locked_slides=[1],
        prior_deck_path=first.deck_path,
    )
    copied = Presentation(second.deck_path).slides[0].shapes[0].text_frame.text
    assert copied == original
    assert "THIS HEADLINE MUST NOT APPEAR" not in copied
    # Unlocked slide still follows the new plan.
    unlocked = Presentation(second.deck_path).slides[1].shapes[0].text_frame.text
    assert "Clinical context" in unlocked
