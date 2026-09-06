"""One-shot builder for clean_deck.pptx and broken_deck.pptx. Not part of the runtime API."""

from __future__ import annotations

import struct
import zlib
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def _png(width: int, height: int, rgb: tuple[int, int, int]) -> bytes:
    raw = b"".join(b"\x00" + (bytes(rgb) * width) for _ in range(height))

    def chunk(tag: bytes, data: bytes) -> bytes:
        return struct.pack(">I", len(data)) + tag + data + struct.pack(">I", zlib.crc32(tag + data) & 0xFFFFFFFF)

    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", width, height, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress(raw, 9))
        + chunk(b"IEND", b"")
    )


def _set_run(paragraph, text: str, *, name: str, size: float, color: RGBColor, bold: bool = False) -> None:
    paragraph.clear()
    run = paragraph.add_run()
    run.text = text
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.bold = bold


def _textbox(slide, left, top, width, height, text, *, name="Arial", size=16, color=None, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    _set_run(tf.paragraphs[0], text, name=name, size=size, color=color or RGBColor(0, 0, 0), bold=bold)
    return box


def build_clean(path: Path, png_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank

    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.22), Inches(12.0), Inches(0.06))
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(0, 51, 102)
    bar.line.fill.background()

    title = _textbox(
        slide,
        Inches(0.6),
        Inches(0.45),
        Inches(12.0),
        Inches(0.7),
        "Demo Trial 2024 — HFrEF outcomes",
        name="Arial",
        size=28,
        color=RGBColor(0, 0, 0),
        bold=True,
    )
    title.text_frame.paragraphs[0].alignment = PP_ALIGN.LEFT

    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.4), Inches(12.0), Inches(2.8))
    tf = body.text_frame
    tf.word_wrap = True
    bullets = [
        "Randomized, double-blind design in adults with HFrEF.",
        "Primary endpoint met — HR 0.78 (95% CI 0.68–0.90).",
        "Serious adverse events in 18.2% of treated patients.",
    ]
    for i, line in enumerate(bullets):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_run(para, line, name="Arial", size=16, color=RGBColor(0, 0, 0))

    slide.shapes.add_picture(str(png_path), Inches(0.6), Inches(5.0), Inches(1.5), Inches(1.5))
    prs.save(str(path))


def build_broken(path: Path, png_path: Path) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # 1+2 text overflow + extra font: tiny box, long 18pt copy
    overflow = (
        "Primary endpoint results require far more copy than this box can hold "
        "and the estimated rendered height will exceed the text frame."
    )
    _textbox(
        slide,
        Inches(0.25),
        Inches(0.2),
        Inches(1.7),
        Inches(0.28),
        overflow,
        name="Arial",
        size=18,
        color=RGBColor(20, 20, 20),
    )

    # 3 min font + 4 family Times + 7 contrast (gray on white)
    _textbox(
        slide,
        Inches(2.2),
        Inches(0.2),
        Inches(2.2),
        Inches(0.7),
        "Tiny gray caption",
        name="Times New Roman",
        size=10,
        color=RGBColor(190, 190, 190),
    )

    # More families + colors
    _textbox(
        slide,
        Inches(4.6),
        Inches(0.2),
        Inches(2.0),
        Inches(0.5),
        "Georgia accent",
        name="Georgia",
        size=16,
        color=RGBColor(128, 0, 128),
    )
    _textbox(
        slide,
        Inches(6.9),
        Inches(0.2),
        Inches(2.0),
        Inches(0.5),
        "Calibri accent",
        name="Calibri",
        size=16,
        color=RGBColor(0, 128, 128),
    )

    # 5 image distortion (100x100 placed 3.2 x 1.0)
    slide.shapes.add_picture(str(png_path), Inches(9.1), Inches(0.15), Inches(3.2), Inches(1.0))

    # 6 shape past right edge + extra color
    overflow_shape = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(12.4), Inches(1.4), Inches(2.0), Inches(0.8)
    )
    overflow_shape.fill.solid()
    overflow_shape.fill.fore_color.rgb = RGBColor(255, 0, 0)
    overflow_shape.line.fill.background()

    # 11 overlap + colors
    a = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.3), Inches(1.6), Inches(2.6), Inches(1.6))
    a.fill.solid()
    a.fill.fore_color.rgb = RGBColor(0, 160, 0)
    a.line.fill.background()
    b = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.4), Inches(2.2), Inches(2.6), Inches(1.6))
    b.fill.solid()
    b.fill.fore_color.rgb = RGBColor(255, 128, 0)
    b.line.fill.background()

    # 8 density: many bullets + many words  (left edge 4.2 — another cluster)
    dense = slide.shapes.add_textbox(Inches(4.2), Inches(1.5), Inches(4.4), Inches(4.8))
    tf = dense.text_frame
    tf.word_wrap = True
    lines = [
        "First extra bullet about secondary findings that should not crowd this slide at all.",
        "Second extra bullet listing exploratory analyses without a clear takeaway sentence.",
        "Third extra bullet repeating safety numbers that belong on a dedicated safety slide.",
        "Fourth extra bullet with population notes that overfill the density budget badly.",
        "Fifth extra bullet describing limitations that make the word count blow past eighty.",
        "Sixth extra bullet adding baseline characteristics nobody can scan in this layout.",
        "Seventh extra bullet guaranteeing this slide exceeds both word and bullet limits.",
        "Eighth extra bullet so Gate 2 density must fail on this deliberately broken deck.",
    ]
    for i, line in enumerate(lines):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        _set_run(para, line, name="Comic Sans MS", size=14, color=RGBColor(0, 0, 180))

    # 10 empty / default placeholder
    _textbox(
        slide,
        Inches(8.9),
        Inches(1.6),
        Inches(3.0),
        Inches(0.6),
        "Click to add title",
        name="Arial",
        size=18,
        color=RGBColor(0, 0, 0),
    )
    empty = slide.shapes.add_textbox(Inches(8.9), Inches(2.4), Inches(3.0), Inches(0.5))
    empty.text_frame.word_wrap = True

    # Extra left-edge cluster to guarantee >4 (0.25, 2.2, 4.6, 6.9, 9.1, 12.4, 0.3, 1.4, 4.2, 8.9)
    prs.save(str(path))


def main() -> None:
    fixtures = Path(__file__).resolve().parents[2] / "tests" / "fixtures"
    fixtures.mkdir(parents=True, exist_ok=True)
    png_path = fixtures / "_pixel.png"
    png_path.write_bytes(_png(100, 100, (30, 90, 150)))
    build_clean(fixtures / "clean_deck.pptx", png_path)
    build_broken(fixtures / "broken_deck.pptx", png_path)
    png_path.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
