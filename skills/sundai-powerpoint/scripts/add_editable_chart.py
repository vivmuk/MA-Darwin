#!/usr/bin/env python3
"""Editable native OOXML clustered-column charts (Excel-backed).

Primary chart path for sundai-powerpoint 3.2.0-darwin.
Users can Edit Data in PowerPoint/Excel after open.

CLI:  python add_editable_chart.py --demo /tmp/demo_surmount_chart.pptx
Lib:  from add_editable_chart import add_clustered_column_chart
"""
from __future__ import annotations

import argparse
import zipfile
from pathlib import Path
from typing import Mapping, Sequence

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

PALETTE = {
    "navy": "12283A",
    "teal": "0E7C7B",
    "muted_teal": "5B9A98",
    "oxblood": "8C2F39",
    "slate": "5B6B79",
    "cloud": "E8ECEF",
    "gold": "C89B3C",
    "white": "FFFFFF",
}

DEFAULT_SERIES_COLORS = (
    PALETTE["oxblood"],
    PALETTE["teal"],
    PALETTE["muted_teal"],
    PALETTE["navy"],
    PALETTE["slate"],
)


def _rgb(hex6: str) -> RGBColor:
    h = hex6.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def build_chart_data(
    categories: Sequence[str],
    series: Mapping[str, Sequence[float]],
) -> CategoryChartData:
    chart_data = CategoryChartData()
    chart_data.categories = list(categories)
    n = len(categories)
    for name, values in series.items():
        vals = list(values)
        if len(vals) != n:
            raise ValueError(f"series {name!r} length {len(vals)} != categories {n}")
        chart_data.add_series(name, tuple(vals))
    return chart_data


def style_clustered_chart(chart, series_colors: Sequence[str] | None = None) -> None:
    colors = list(series_colors or DEFAULT_SERIES_COLORS)
    chart.plots[0].has_data_labels = False
    for i, ser in enumerate(chart.series):
        hex6 = colors[i % len(colors)]
        try:
            ser.format.fill.solid()
            ser.format.fill.fore_color.rgb = _rgb(hex6)
        except Exception:
            pass
    chart.has_legend = True
    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
    chart.legend.include_in_layout = False
    try:
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = _rgb(PALETTE["cloud"])
    except Exception:
        pass
    try:
        chart.category_axis.tick_labels.font.size = Pt(9)
        chart.category_axis.tick_labels.font.color.rgb = _rgb(PALETTE["slate"])
        chart.value_axis.tick_labels.font.size = Pt(9)
        chart.value_axis.tick_labels.font.color.rgb = _rgb(PALETTE["slate"])
    except Exception:
        pass


def add_clustered_column_chart(
    slide,
    categories: Sequence[str],
    series: Mapping[str, Sequence[float]],
    left=Inches(0.5),
    top=Inches(1.8),
    width=Inches(8.5),
    height=Inches(4.2),
    series_colors: Sequence[str] | None = None,
):
    """Add editable clustered-column chart (ppt/charts + ppt/embeddings xlsx)."""
    chart_data = build_chart_data(categories, series)
    graphic_frame = slide.shapes.add_chart(
        XL_CHART_TYPE.COLUMN_CLUSTERED, left, top, width, height, chart_data
    )
    chart = graphic_frame.chart
    style_clustered_chart(chart, series_colors=series_colors)
    return chart


SURMOUNT_CATEGORIES = [
    "ASCVD (ACC/AHA)",
    "ASCVD (PREVENT)",
    "HF (PREVENT)",
    "Total CVD (PREVENT)",
]
SURMOUNT_SERIES = {
    "Placebo": (57.9, 40.5, 56.3, 43.8),
    "TZP 5": (-4.6, -3.7, -0.7, 1.3),
    "TZP 10": (-7.5, -6.3, 1.6, -0.7),
    "TZP 15": (-9.2, -8.8, -5.4, -2.4),
}


def write_demo_pptx(out_path: Path) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    rule = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.333), Inches(0.08))
    rule.fill.solid()
    rule.fill.fore_color.rgb = _rgb(PALETTE["teal"])
    rule.line.fill.background()

    title = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12), Inches(0.5))
    p = title.text_frame.paragraphs[0]
    p.text = "THE EVIDENCE — PRIMARY RESULT (demo editable chart)"
    p.font.name = "Arial"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = _rgb(PALETTE["navy"])

    add_clustered_column_chart(
        slide, SURMOUNT_CATEGORIES, SURMOUNT_SERIES,
        left=Inches(0.5), top=Inches(1.1), width=Inches(8.8), height=Inches(5.2),
    )

    callout = slide.shapes.add_shape(1, Inches(9.5), Inches(1.5), Inches(3.4), Inches(3.2))
    callout.fill.solid()
    callout.fill.fore_color.rgb = _rgb(PALETTE["white"])
    callout.line.color.rgb = _rgb(PALETTE["gold"])
    callout.line.width = Pt(1.5)
    ctf = callout.text_frame
    ctf.word_wrap = True
    cp = ctf.paragraphs[0]
    cp.text = (
        "Absolute median callout (example):\n"
        "Placebo ASCVD ~2.4% → 3.7%\n"
        "Large % change on low baseline — do not over-read +57.9% alone.\n"
        "CIs not reported for LS means in source panel."
    )
    cp.font.name = "Arial"
    cp.font.size = Pt(11)
    cp.font.color.rgb = _rgb(PALETTE["navy"])

    slide.notes_slide.notes_text_frame.text = (
        "Demo: native OOXML clustered column with embedded Excel. "
        "Open Edit Data in PowerPoint to change series. "
        "Increases (HF TZP 10 +1.6%, Total CVD TZP 5 +1.3%) shown honestly."
    )

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return out_path


def main() -> None:
    parser = argparse.ArgumentParser(description="Editable OOXML clustered-column charts.")
    parser.add_argument("--demo", nargs="?", const="/tmp/demo_surmount_chart.pptx", default=None)
    parser.add_argument("--xlsx", default=None)
    args = parser.parse_args()

    if args.xlsx or args.demo:
        try:
            from xlsxwriter import Workbook
            xpath = Path(args.xlsx or "/tmp/demo_surmount_chart_data.xlsx")
            wb = Workbook(str(xpath))
            ws = wb.add_worksheet("LS_mean_pct")
            ws.write_row(0, 0, ["Endpoint"] + list(SURMOUNT_SERIES.keys()))
            for r, cat in enumerate(SURMOUNT_CATEGORIES, start=1):
                ws.write_row(r, 0, [cat] + [SURMOUNT_SERIES[k][r - 1] for k in SURMOUNT_SERIES])
            wb.close()
            print(f"wrote companion xlsx: {xpath}")
        except Exception as exc:
            print(f"companion xlsx skipped: {exc}")

    if args.demo is not None:
        path = write_demo_pptx(Path(args.demo))
        print(f"wrote demo pptx: {path}")
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            print("ppt/charts:", [n for n in names if n.startswith("ppt/charts/")])
            print("ppt/embeddings:", [n for n in names if n.startswith("ppt/embeddings/")])
        return
    parser.print_help()


if __name__ == "__main__":
    main()
