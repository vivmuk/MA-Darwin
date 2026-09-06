"""Gate 2 — visual auto-checks (PRD §7.7 / Prompt 3).

All checks read pptx XML via python-pptx (and rendered PNGs for contrast
fallback). No LLM calls.
"""

from __future__ import annotations

import math
import re
from collections.abc import Iterable, Iterator
from pathlib import Path
from typing import Optional

import yaml
from pptx import Presentation
from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.slide import Slide

from app.models.common import CheckLocation
from app.models.gates import Gate2Result, GateCheckResult
from app.paths import CONFIG_DIR

EMU_PER_PT = 12700.0
BOUNDS_SLACK_EMU = 20000  # ~1.6pt — snapping / rounding, not a real overflow
OVERLAP_MIN_AREA_SQPT = 4.0
ALIGN_TOL_PT = 7.2  # 0.1 inch
LINE_HEIGHT_FACTOR = 1.2
AVG_CHAR_WIDTH_FACTOR = 0.55

DEFAULT_PLACEHOLDER_TEXTS = {
    "click to add title",
    "click to add subtitle",
    "click to add text",
    "click to add notes",
    "click to add content",
}

CHECK_NAMES = (
    "shape_bounds",
    "text_overflow",
    "min_font_size",
    "font_family_count",
    "color_count",
    "image_aspect_distortion",
    "contrast_ratio",
    "alignment_discipline",
    "density",
    "empty_placeholders",
    "shape_overlap",
)


def location(
    slide: int,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    element_id: str | None = None,
    detail: str = "",
) -> CheckLocation:
    """Build a typed Gate 2 failure location."""
    return CheckLocation(
        slide=slide,
        element_id=element_id,
        x=x,
        y=y,
        w=w,
        h=h,
        detail=detail,
    )


def _emu_pt(emu: int | float | None) -> float:
    if emu is None:
        return 0.0
    return float(emu) / EMU_PER_PT


def _result(name: str, locations: list[CheckLocation], ok_msg: str, fail_msg: str) -> GateCheckResult:
    passed = not locations
    return GateCheckResult(
        name=name,
        passed=passed,
        message=ok_msg if passed else fail_msg,
        flagged=False,
        locations=locations,
    )


def _iter_shapes(slide: Slide) -> Iterator[object]:
    for shape in slide.shapes:
        yield shape
        if getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.GROUP:
            try:
                for child in shape.shapes:
                    yield child
            except (AttributeError, ValueError):
                continue


def _shape_id(shape: object) -> str:
    name = getattr(shape, "name", None) or ""
    sid = getattr(shape, "shape_id", None)
    if name and sid is not None:
        return f"{name}:{sid}"
    if name:
        return str(name)
    return f"shape_{sid}"


def _box(shape: object) -> tuple[float, float, float, float]:
    return (
        _emu_pt(getattr(shape, "left", 0) or 0),
        _emu_pt(getattr(shape, "top", 0) or 0),
        _emu_pt(getattr(shape, "width", 0) or 0),
        _emu_pt(getattr(shape, "height", 0) or 0),
    )


def _loc(slide_n: int, shape: object, detail: str = "") -> CheckLocation:
    x, y, w, h = _box(shape)
    return location(slide_n, x, y, w, h, element_id=_shape_id(shape), detail=detail)


def _has_text_frame(shape: object) -> bool:
    try:
        return bool(getattr(shape, "has_text_frame", False))
    except ValueError:
        return False


def _run_size_pt(run: object) -> Optional[float]:
    font = getattr(run, "font", None)
    size = getattr(font, "size", None) if font is not None else None
    if size is not None:
        return float(size.pt)
    r = getattr(run, "_r", None)
    if r is None:
        return None
    r_pr = r.find(qn("a:rPr"))
    if r_pr is not None and r_pr.get("sz"):
        try:
            return int(r_pr.get("sz")) / 100.0
        except ValueError:
            return None
    return None


def _run_font_name(run: object) -> Optional[str]:
    font = getattr(run, "font", None)
    name = getattr(font, "name", None) if font is not None else None
    if name:
        return str(name)
    r = getattr(run, "_r", None)
    if r is None:
        return None
    r_pr = r.find(qn("a:rPr"))
    if r_pr is None:
        return None
    latin = r_pr.find(qn("a:latin"))
    if latin is not None and latin.get("typeface"):
        typeface = latin.get("typeface")
        if typeface and not typeface.startswith("+"):
            return typeface
    return None


def _rgb_tuple_from_color(color: object) -> Optional[tuple[int, int, int]]:
    if color is None:
        return None
    try:
        rgb = getattr(color, "rgb", None)
    except AttributeError:
        rgb = None
    if rgb is None:
        return None
    try:
        return (int(rgb[0]), int(rgb[1]), int(rgb[2]))
    except (TypeError, ValueError, IndexError):
        try:
            hex_s = str(rgb)
            if len(hex_s) == 6:
                return (int(hex_s[0:2], 16), int(hex_s[2:4], 16), int(hex_s[4:6], 16))
        except ValueError:
            return None
    return None


def _run_rgb(run: object) -> Optional[tuple[int, int, int]]:
    font = getattr(run, "font", None)
    if font is None:
        return None
    try:
        return _rgb_tuple_from_color(font.color)
    except AttributeError:
        return None


def _solid_fill_rgb(shape: object) -> Optional[tuple[int, int, int]]:
    try:
        fill = shape.fill
    except (AttributeError, ValueError, TypeError):
        return None
    try:
        if fill.type != MSO_FILL.SOLID:
            return None
        return _rgb_tuple_from_color(fill.fore_color)
    except (AttributeError, TypeError, ValueError):
        return None


def _shape_text(shape: object) -> str:
    if not _has_text_frame(shape):
        return ""
    try:
        return shape.text_frame.text or ""
    except (AttributeError, ValueError):
        return ""


def _iter_runs(shape: object) -> Iterator[object]:
    if not _has_text_frame(shape):
        return
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            yield run


def _load_defaults(defaults_config: Path | str | None) -> dict[str, object]:
    path = Path(defaults_config) if defaults_config is not None else CONFIG_DIR / "defaults.yaml"
    with path.open(encoding="utf-8") as fh:
        data = yaml.safe_load(fh) or {}
    if not isinstance(data, dict):
        raise ValueError(f"defaults config must be a mapping: {path}")
    return data


def check_shape_bounds(pptx_path: Path | str) -> GateCheckResult:
    """No shape extends past slide edges."""
    prs = Presentation(str(pptx_path))
    slide_w = float(prs.slide_width)
    slide_h = float(prs.slide_height)
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            left = float(getattr(shape, "left", 0) or 0)
            top = float(getattr(shape, "top", 0) or 0)
            width = float(getattr(shape, "width", 0) or 0)
            height = float(getattr(shape, "height", 0) or 0)
            overflow = (
                left < -BOUNDS_SLACK_EMU
                or top < -BOUNDS_SLACK_EMU
                or left + width > slide_w + BOUNDS_SLACK_EMU
                or top + height > slide_h + BOUNDS_SLACK_EMU
            )
            if overflow:
                locations.append(_loc(i, shape, "shape extends past slide edge"))
    return _result(
        "shape_bounds",
        locations,
        "All shapes within slide edges",
        f"{len(locations)} shape(s) extend past slide edges",
    )


def _estimate_text_height_pt(shape: object) -> float:
    if not _has_text_frame(shape):
        return 0.0
    box_w = max(_emu_pt(getattr(shape, "width", 0) or 0), 1.0)
    used = 0.0
    for paragraph in shape.text_frame.paragraphs:
        runs = list(paragraph.runs)
        text = paragraph.text or ""
        font_pt = 14.0
        for run in runs:
            sz = _run_size_pt(run)
            if sz is not None:
                font_pt = sz
                break
        if not text.strip():
            used += font_pt * 0.35
            continue
        avg_w = max(font_pt * AVG_CHAR_WIDTH_FACTOR, 1.0)
        chars_per_line = max(1, int(box_w / avg_w))
        lines = 0
        for part in text.splitlines() or [text]:
            lines += max(1, math.ceil(len(part) / chars_per_line))
        used += lines * font_pt * LINE_HEIGHT_FACTOR
    return used


def check_text_overflow(pptx_path: Path | str) -> GateCheckResult:
    """Estimated rendered text height must fit the text box height."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            if not _has_text_frame(shape):
                continue
            text = _shape_text(shape).strip()
            if not text:
                continue
            box_h = _emu_pt(getattr(shape, "height", 0) or 0)
            if box_h <= 0:
                continue
            estimated = _estimate_text_height_pt(shape)
            if estimated > box_h * 1.05:
                locations.append(
                    _loc(
                        i,
                        shape,
                        f"estimated text height {estimated:.1f}pt > box {box_h:.1f}pt",
                    )
                )
    return _result(
        "text_overflow",
        locations,
        "No text overflow detected",
        f"{len(locations)} text box(es) overflow their bounds",
    )


def check_min_font_size(pptx_path: Path | str, *, min_font_pt: float) -> GateCheckResult:
    """No text run below configured minimum body size."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            for run in _iter_runs(shape):
                if not (getattr(run, "text", "") or "").strip():
                    continue
                size = _run_size_pt(run)
                if size is not None and size + 1e-6 < float(min_font_pt):
                    locations.append(_loc(i, shape, f"{size:g}pt run"))
                    break
    return _result(
        "min_font_size",
        locations,
        f"All text runs are ≥ {min_font_pt:g}pt",
        f"Text below {min_font_pt:g}pt on {len(locations)} shape(s)",
    )


def check_font_family_count(pptx_path: Path | str, *, max_families: int = 2) -> GateCheckResult:
    """Distinct font family count must be ≤ ``max_families``."""
    prs = Presentation(str(pptx_path))
    families: set[str] = set()
    sample: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            for run in _iter_runs(shape):
                name = _run_font_name(run)
                if not name:
                    continue
                key = name.strip()
                if key.lower() not in {f.lower() for f in families}:
                    families.add(key)
                    if len(families) > max_families:
                        sample.append(_loc(i, shape, f"family {key}"))
    extra = sorted(families)
    locations = sample if len(families) > max_families else []
    return _result(
        "font_family_count",
        locations,
        f"{len(families)} font family(ies) ≤ {max_families}",
        f"{len(families)} font families ({', '.join(extra)}) exceed max {max_families}",
    )


def _collect_colors(prs: Presentation) -> tuple[set[tuple[int, int, int]], list[CheckLocation]]:
    colors: set[tuple[int, int, int]] = set()
    samples: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            fill = _solid_fill_rgb(shape)
            if fill is not None and fill not in colors:
                colors.add(fill)
                samples.append(_loc(i, shape, f"fill {fill}"))
            for run in _iter_runs(shape):
                rgb = _run_rgb(run)
                if rgb is not None and rgb not in colors:
                    colors.add(rgb)
                    samples.append(_loc(i, shape, f"font {rgb}"))
    return colors, samples


def check_color_count(pptx_path: Path | str, *, max_colors: int = 5) -> GateCheckResult:
    """Distinct colour count must be ≤ ``max_colors``."""
    prs = Presentation(str(pptx_path))
    colors, samples = _collect_colors(prs)
    locations = samples[max_colors:] if len(colors) > max_colors else []
    return _result(
        "color_count",
        locations,
        f"{len(colors)} colour(s) ≤ {max_colors}",
        f"{len(colors)} colours exceed max {max_colors}",
    )


def _png_size(blob: bytes) -> Optional[tuple[int, int]]:
    if len(blob) >= 24 and blob[:8] == b"\x89PNG\r\n\x1a\n":
        w = int.from_bytes(blob[16:20], "big")
        h = int.from_bytes(blob[20:24], "big")
        if w > 0 and h > 0:
            return w, h
    return None


def check_image_aspect_distortion(
    pptx_path: Path | str,
    *,
    tolerance: float = 0.02,
) -> GateCheckResult:
    """Placed image w/h vs native ratio within tolerance (default 2%)."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            if getattr(shape, "shape_type", None) != MSO_SHAPE_TYPE.PICTURE:
                continue
            try:
                image = shape.image
                native = image.size
            except Exception:
                native = None
            if native is None:
                try:
                    native = _png_size(shape.image.blob)
                except Exception:
                    native = None
            if not native or native[1] == 0:
                continue
            placed_w = float(getattr(shape, "width", 0) or 0)
            placed_h = float(getattr(shape, "height", 0) or 0)
            if placed_h == 0:
                continue
            placed_ratio = placed_w / placed_h
            native_ratio = native[0] / native[1]
            rel = abs(placed_ratio - native_ratio) / native_ratio
            if rel > tolerance:
                locations.append(
                    _loc(
                        i,
                        shape,
                        f"aspect distortion {rel:.1%} (tolerance {tolerance:.1%})",
                    )
                )
    return _result(
        "image_aspect_distortion",
        locations,
        "No image aspect distortion",
        f"{len(locations)} image(s) stretched beyond {tolerance:.0%} tolerance",
    )


def _linearize(channel: float) -> float:
    c = channel / 255.0
    if c <= 0.04045:
        return c / 12.92
    return ((c + 0.055) / 1.055) ** 2.4


def wcag_contrast(fg: tuple[int, int, int], bg: tuple[int, int, int]) -> float:
    """WCAG 2.x contrast ratio between two sRGB colours."""
    def lum(rgb: tuple[int, int, int]) -> float:
        r, g, b = rgb
        return 0.2126 * _linearize(r) + 0.7152 * _linearize(g) + 0.0722 * _linearize(b)

    l1, l2 = lum(fg), lum(bg)
    lighter, darker = (l1, l2) if l1 >= l2 else (l2, l1)
    return (lighter + 0.05) / (darker + 0.05)


def _sample_png_bg(
    image_path: Path,
    shape: object,
    slide_w_pt: float,
    slide_h_pt: float,
) -> Optional[tuple[int, int, int]]:
    try:
        import pymupdf
    except ImportError:
        return None
    x, y, w, h = _box(shape)
    try:
        pix = pymupdf.Pixmap(str(image_path))
    except Exception:
        return None
    try:
        if pix.n < 3 or slide_w_pt <= 0 or slide_h_pt <= 0:
            return None
        # Sample a pixel just inside the box, slightly above the first text line.
        sx = int(max(0, min(pix.width - 1, (x + w * 0.05) / slide_w_pt * pix.width)))
        sy = int(max(0, min(pix.height - 1, (y + 2.0) / slide_h_pt * pix.height)))
        pixel = pix.pixel(sx, sy)
        return (int(pixel[0]), int(pixel[1]), int(pixel[2]))
    except Exception:
        return None
    finally:
        pix = None


def check_contrast_ratio(
    pptx_path: Path | str,
    slide_images: list[Path | str],
    *,
    min_contrast: float,
) -> GateCheckResult:
    """Text vs background fill contrast must be ≥ configured WCAG ratio."""
    prs = Presentation(str(pptx_path))
    images = [Path(p) for p in slide_images]
    slide_w_pt = _emu_pt(prs.slide_width)
    slide_h_pt = _emu_pt(prs.slide_height)
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        png = images[i - 1] if i - 1 < len(images) else None
        for shape in _iter_shapes(slide):
            if not _has_text_frame(shape) or not _shape_text(shape).strip():
                continue
            bg = _solid_fill_rgb(shape)
            if bg is None and png is not None and png.is_file():
                bg = _sample_png_bg(png, shape, slide_w_pt, slide_h_pt)
            if bg is None:
                bg = (255, 255, 255)
            worst: Optional[float] = None
            for run in _iter_runs(shape):
                if not (getattr(run, "text", "") or "").strip():
                    continue
                fg = _run_rgb(run) or (0, 0, 0)
                ratio = wcag_contrast(fg, bg)
                if worst is None or ratio < worst:
                    worst = ratio
            if worst is not None and worst + 1e-6 < float(min_contrast):
                locations.append(_loc(i, shape, f"contrast {worst:.2f}:1 < {min_contrast}:1"))
    return _result(
        "contrast_ratio",
        locations,
        f"All text meets {min_contrast}:1",
        f"{len(locations)} text region(s) below {min_contrast}:1 contrast",
    )


def _cluster_count(values: Iterable[float], tol: float) -> int:
    ordered = sorted(values)
    if not ordered:
        return 0
    clusters = 1
    start = ordered[0]
    for value in ordered[1:]:
        if value - start > tol:
            clusters += 1
            start = value
    return clusters


def check_alignment_discipline(pptx_path: Path | str) -> GateCheckResult:
    """Left edges must cluster to ≤ 4 distinct x values per slide."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        xs: list[float] = []
        shapes = list(_iter_shapes(slide))
        for shape in shapes:
            left = _emu_pt(getattr(shape, "left", 0) or 0)
            width = _emu_pt(getattr(shape, "width", 0) or 0)
            height = _emu_pt(getattr(shape, "height", 0) or 0)
            if width <= 0 or height <= 0:
                continue
            xs.append(left)
        clusters = _cluster_count(xs, ALIGN_TOL_PT)
        if clusters > 4:
            # Pin the first shape as the slide-level finding.
            if shapes:
                locations.append(_loc(i, shapes[0], f"{clusters} left-edge clusters"))
            else:
                locations.append(location(i, 0, 0, 0, 0, detail=f"{clusters} left-edge clusters"))
    return _result(
        "alignment_discipline",
        locations,
        "Left edges snap to ≤ 4 x-values per slide",
        f"{len(locations)} slide(s) have more than 4 left-edge clusters",
    )


def check_density(
    pptx_path: Path | str,
    *,
    max_words_per_slide: int,
    max_bullets_per_slide: int,
) -> GateCheckResult:
    """Word and bullet counts per slide vs skill-defined limits."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    word_re = re.compile(r"\S+")
    for i, slide in enumerate(prs.slides, start=1):
        words = 0
        bullets = 0
        first_shape = None
        for shape in _iter_shapes(slide):
            if not _has_text_frame(shape):
                continue
            if first_shape is None:
                first_shape = shape
            for paragraph in shape.text_frame.paragraphs:
                text = (paragraph.text or "").strip()
                if not text:
                    continue
                words += len(word_re.findall(text))
                bullets += 1
        if words > max_words_per_slide or bullets > max_bullets_per_slide:
            target = first_shape
            detail = f"{words} words, {bullets} bullets"
            if target is not None:
                locations.append(_loc(i, target, detail))
            else:
                locations.append(location(i, 0, 0, 0, 0, detail=detail))
    return _result(
        "density",
        locations,
        "Word and bullet counts within limits",
        f"{len(locations)} slide(s) exceed density limits",
    )


def check_empty_placeholders(pptx_path: Path | str) -> GateCheckResult:
    """Fail on empty or default-text placeholder elements."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        for shape in _iter_shapes(slide):
            is_ph = bool(getattr(shape, "is_placeholder", False))
            is_textbox = getattr(shape, "shape_type", None) == MSO_SHAPE_TYPE.TEXT_BOX
            if not _has_text_frame(shape):
                continue
            text = _shape_text(shape).strip()
            default = text.lower() in DEFAULT_PLACEHOLDER_TEXTS
            empty = text == ""
            if default or (empty and (is_ph or is_textbox)):
                detail = "default placeholder text" if default else "empty placeholder"
                locations.append(_loc(i, shape, detail))
    return _result(
        "empty_placeholders",
        locations,
        "No empty or default-text placeholders",
        f"{len(locations)} empty or placeholder-looking element(s)",
    )


def _contains(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> bool:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    return ax <= bx + 0.5 and ay <= by + 0.5 and ax + aw >= bx + bw - 0.5 and ay + ah >= by + bh - 0.5


def _intersection_area(a: tuple[float, float, float, float], b: tuple[float, float, float, float]) -> float:
    ax, ay, aw, ah = a
    bx, by, bw, bh = b
    x1 = max(ax, bx)
    y1 = max(ay, by)
    x2 = min(ax + aw, bx + bw)
    y2 = min(ay + ah, by + bh)
    if x2 <= x1 or y2 <= y1:
        return 0.0
    return (x2 - x1) * (y2 - y1)


def check_shape_overlap(pptx_path: Path | str) -> GateCheckResult:
    """Bounding-box intersection between non-nested / non-intentional shapes."""
    prs = Presentation(str(pptx_path))
    locations: list[CheckLocation] = []
    for i, slide in enumerate(prs.slides, start=1):
        boxes: list[tuple[object, tuple[float, float, float, float]]] = []
        for shape in _iter_shapes(slide):
            box = _box(shape)
            if box[2] <= 0 or box[3] <= 0:
                continue
            # Skip very thin lines / accent rules from overlap (intentional bars).
            if box[3] < 6.0 or box[2] < 6.0:
                continue
            boxes.append((shape, box))
        for idx, (shape_a, box_a) in enumerate(boxes):
            for shape_b, box_b in boxes[idx + 1 :]:
                if _contains(box_a, box_b) or _contains(box_b, box_a):
                    continue
                area = _intersection_area(box_a, box_b)
                if area > OVERLAP_MIN_AREA_SQPT:
                    locations.append(
                        _loc(i, shape_a, f"overlaps {_shape_id(shape_b)} ({area:.1f} sq pt)")
                    )
                    break
    return _result(
        "shape_overlap",
        locations,
        "No unintentional shape overlaps",
        f"{len(locations)} unintentional overlapping shape pair(s)",
    )


def run_gate2(
    *,
    pptx_path: Path | str,
    slide_images: list[Path | str],
    defaults_config: Path | str | None = None,
) -> Gate2Result:
    """Run all Gate 2 checks. No LLM calls. Failures block Gate 3."""
    cfg = _load_defaults(defaults_config)
    checks = [
        check_shape_bounds(pptx_path),
        check_text_overflow(pptx_path),
        check_min_font_size(pptx_path, min_font_pt=float(cfg["min_font_pt"])),
        check_font_family_count(pptx_path, max_families=int(cfg["max_font_families"])),
        check_color_count(pptx_path, max_colors=int(cfg["max_colors"])),
        check_image_aspect_distortion(
            pptx_path,
            tolerance=float(cfg.get("image_aspect_tolerance", 0.02)),
        ),
        check_contrast_ratio(
            pptx_path,
            slide_images,
            min_contrast=float(cfg["min_contrast"]),
        ),
        check_alignment_discipline(pptx_path),
        check_density(
            pptx_path,
            max_words_per_slide=int(cfg["max_words_per_slide"]),
            max_bullets_per_slide=int(cfg["max_bullets_per_slide"]),
        ),
        check_empty_placeholders(pptx_path),
        check_shape_overlap(pptx_path),
    ]
    return Gate2Result(gate="gate2", passed=all(c.passed for c in checks), checks=checks)
