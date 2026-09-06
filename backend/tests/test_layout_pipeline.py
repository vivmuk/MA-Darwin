"""LayoutSpec is the source of truth: SVG, pptx, and CLI compare-layout."""

from __future__ import annotations

from pathlib import Path

from app.cli import main
from app.generation.pptx_writer import write_pptx
from app.models.layout import FontWeight, LayoutElement, LayoutElementKind, LayoutSlide, LayoutSpec
from app.rendering.svg_renderer import write_slide_svgs
from app.rendering.text_metrics import text_exceeds_box


def _spec() -> LayoutSpec:
    return LayoutSpec(
        slides=[
            LayoutSlide(
                slide=1,
                role="title",
                elements=[
                    LayoutElement(
                        id="title",
                        kind=LayoutElementKind.TEXT,
                        x=0.5,
                        y=0.4,
                        w=12.0,
                        h=1.0,
                        font_family="Arial",
                        font_size_pt=28,
                        font_weight=FontWeight.BOLD,
                        color="#1B2A4A",
                        text="Primary endpoint met",
                    )
                ],
            )
        ]
    )


def test_text_exceeds_box_detects_overflow() -> None:
    tight = LayoutElement(
        id="tiny",
        x=0.2,
        y=0.2,
        w=1.2,
        h=0.25,
        font_size_pt=18,
        text="Primary endpoint results require far more copy than this box can hold",
    )
    assert text_exceeds_box(tight)
    roomy = tight.model_copy(update={"w": 12.0, "h": 1.2})
    assert text_exceeds_box(roomy) is None


def test_compare_layout_cli_writes_svg_and_pptx(tmp_path: Path) -> None:
    layout = tmp_path / "layout_spec.json"
    layout.write_text(_spec().model_dump_json(), encoding="utf-8")
    out = tmp_path / "side_by_side"
    code = main(["compare-layout", "--layout", str(layout), "--out", str(out)])
    assert code == 0
    assert (out / "svg" / "slide_01.svg").is_file()
    assert "Primary endpoint met" in (out / "svg" / "slide_01.svg").read_text(encoding="utf-8")
    assert (out / "deck.pptx").is_file()


def test_pptx_and_svg_share_coordinates(tmp_path: Path) -> None:
    spec = _spec()
    svg = write_slide_svgs(spec, tmp_path / "svg")[0].read_text(encoding="utf-8")
    pptx = write_pptx(spec, tmp_path / "deck.pptx")
    assert 'font-size="28' in svg
    assert 'x="36.0"' in svg  # 0.5 in * 72
    from pptx import Presentation
    from pptx.enum.text import MSO_AUTO_SIZE
    from pptx.util import Emu, Inches

    prs = Presentation(str(pptx))
    box = prs.slides[0].shapes[0]
    assert abs(box.left - Inches(0.5)) < Emu(1000)
    assert box.text_frame.auto_size in (None, MSO_AUTO_SIZE.NONE)
