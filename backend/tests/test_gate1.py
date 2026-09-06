"""Gate 1 deterministic checks (PRD §7.6) against frozen fixtures."""

from __future__ import annotations

import json
from pathlib import Path

import yaml
from pptx import Presentation

from app.gates.gate1_content import (
    check_endpoint_labelling,
    check_number_sweep,
    check_promotional_language,
    run_gate1,
)
from app.generation.generator import generate_deck
from app.models.blueprint import Blueprint
from app.models.claim import ClaimLedger, ClaimLedgerEntry, ClaimType, EvidenceClass
from app.models.numbers import NumbersIndex
from app.models.run import Brief
from app.models.slide import SlideMap, SlideMapEntry
from app.paths import BLUEPRINTS_DIR, CONFIG_DIR
from tests.test_generator import complete_plan_from_fixtures

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def _blueprint() -> Blueprint:
    payload = yaml.safe_load((BLUEPRINTS_DIR / "msl_physician_8.yaml").read_text(encoding="utf-8"))
    return Blueprint.model_validate(payload)


def _ledger() -> ClaimLedger:
    return ClaimLedger.model_validate(json.loads((FIXTURES / "ledger.json").read_text(encoding="utf-8")))


def _numbers() -> NumbersIndex:
    return NumbersIndex.model_validate(json.loads((FIXTURES / "numbers_index.json").read_text(encoding="utf-8")))


def _roles() -> list[str]:
    return [s.role for s in _blueprint().slides]


def _passing_deck(tmp_path: Path):
    result = generate_deck(
        blueprint=_blueprint(),
        ledger=_ledger(),
        brief=Brief(),
        skill_version="v1",
        slide_plan=complete_plan_from_fixtures(),
        output_dir=tmp_path,
    )
    return result


def test_generated_deck_passes_gate1(tmp_path: Path) -> None:
    result = _passing_deck(tmp_path)
    report = run_gate1(
        pptx_path=result.deck_path,
        slide_map=result.slide_map,
        ledger=_ledger(),
        numbers_index=_numbers(),
        blueprint_roles=_roles(),
    )
    failed = [c.name for c in report.checks if not c.passed]
    assert failed == [], {c.name: c.message for c in report.checks if not c.passed}
    assert report.passed is True
    names = {c.name for c in report.checks}
    assert {
        "claim_mapping",
        "number_sweep",
        "endpoint_labelling",
        "fair_balance",
        "comparative_claims",
        "promotional_language",
        "references_complete",
        "synthesized_flags",
    } <= names


def test_number_sweep_fails_on_injected_fake_number(tmp_path: Path) -> None:
    result = _passing_deck(tmp_path / "good")
    mutated = tmp_path / "fake.pptx"
    prs = Presentation(result.deck_path)
    target = next(s for s in prs.slides[4].shapes if getattr(s, "has_text_frame", False))
    target.name = "injected_fake_number"
    para = target.text_frame.paragraphs[0]
    if para.runs:
        para.runs[0].text = f"{para.runs[0].text} N=99999"
    else:
        para.text = f"{target.text_frame.text} N=99999"
    prs.save(str(mutated))

    cfg = yaml.safe_load((CONFIG_DIR / "compliance.yaml").read_text(encoding="utf-8"))
    ns = cfg["number_sweep"]
    check = check_number_sweep(
        mutated,
        _numbers(),
        relative_tolerance=float(ns["relative_tolerance"]),
        absolute_tolerance=float(ns["absolute_tolerance"]),
    )
    assert check.passed is False
    assert check.name == "number_sweep"
    assert any("99999" in loc.detail for loc in check.locations)


def test_endpoint_labelling_fails_when_label_stripped() -> None:
    ledger = _ledger()
    ledger.entries.append(
        ClaimLedgerEntry(
            id="C-030",
            text="Hospitalization was reduced in the treated arm.",
            verbatim="Hospitalization was reduced in the treated arm.",
            page=5,
            section="Results",
            claim_type=ClaimType.PARAPHRASE,
            evidence_class=EvidenceClass.SECONDARY_ENDPOINT,
            numbers=[],
        )
    )
    slide_map = SlideMap(
        deck_path="round_1/deck.pptx",
        entries=[
            SlideMapEntry(
                slide=6,
                element_id="s6_body_0",
                text="Hospitalization was reduced in the treated arm.",
                claim_ids=["C-030"],
            )
        ],
    )
    check = check_endpoint_labelling(slide_map, ledger)
    assert check.passed is False
    assert check.name == "endpoint_labelling"
    assert any("C-030" in loc.detail for loc in check.locations)


def test_promotional_language_fails_on_injected_word() -> None:
    result_map = SlideMap(
        deck_path="round_1/deck.pptx",
        entries=[
            SlideMapEntry(
                slide=1,
                element_id="s1_headline",
                text="Proven HFrEF outcomes",
                claim_ids=["C-001"],
            )
        ],
    )
    cfg = yaml.safe_load((CONFIG_DIR / "compliance.yaml").read_text(encoding="utf-8"))
    check = check_promotional_language(result_map, list(cfg["promotional_blocklist"]))
    assert check.passed is False
    assert check.name == "promotional_language"
    assert any("proven" in loc.detail.lower() for loc in check.locations)
