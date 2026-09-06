"""LayoutSpec → deck.pptx. Absolute inches, explicit fonts, never autofit."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

from app.models.layout import FontWeight, LayoutElement, LayoutElementKind, LayoutSlide, LayoutSpec, TextAlign

_ALIGN = {
    TextAlign.LEFT: PP_ALIGN.LEFT,
    TextAlign.CENTER: PP_ALIGN.CENTER,
    TextAlign.RIGHT: PP_ALIGN.RIGHT,
}


def write_pptx(spec: LayoutSpec, output_path: Path | str) -> Path:
    """Write ``deck.pptx`` with 1:1 coordinates from ``spec``."""
    dest = Path(output_path)
    dest.parent.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = Inches(spec.canvas_width)
    prs.slide_height = Inches(spec.canvas_height)
    for slide_spec in spec.slides:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _apply_slide(slide, slide_spec)
    prs.save(str(dest))
    return dest


def _apply_slide(slide, spec: LayoutSlide) -> None:
    for el in spec.elements:
        if el.kind == LayoutElementKind.IMAGE and el.image_path and Path(el.image_path).is_file():
            pic = slide.shapes.add_picture(
                el.image_path, Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h)
            )
            pic.name = el.id
            continue
        if el.kind == LayoutElementKind.SHAPE:
            shape = slide.shapes.add_shape(
                MSO_SHAPE.RECTANGLE, Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h)
            )
            shape.name = el.id
            fill = el.fill or el.color
            shape.fill.solid()
            shape.fill.fore_color.rgb = _rgb(fill)
            shape.line.fill.background()
            continue
        _textbox(slide, el)
    if spec.speaker_notes:
        slide.notes_slide.notes_text_frame.text = spec.speaker_notes


def _textbox(slide, el: LayoutElement) -> None:
    box = slide.shapes.add_textbox(Inches(el.x), Inches(el.y), Inches(el.w), Inches(el.h))
    box.name = el.id
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    _disable_autofit(tf)
    try:
        tf.vertical_anchor = MSO_ANCHOR.TOP
    except Exception:
        pass
    para = tf.paragraphs[0]
    para.alignment = _ALIGN.get(el.alignment, PP_ALIGN.LEFT)
    run = para.add_run()
    run.text = el.text
    run.font.size = Pt(el.font_size_pt)
    run.font.bold = el.font_weight == FontWeight.BOLD
    run.font.name = el.font_family
    run.font.color.rgb = _rgb(el.color)


def _disable_autofit(text_frame) -> None:
    """Force ``<a:noAutofit/>`` so PowerPoint cannot shrink text."""
    body = text_frame._txBody
    pr = body.find(qn("a:bodyPr"))
    if pr is None:
        return
    for tag in ("a:spAutoFit", "a:normAutofit", "a:noAutofit"):
        node = pr.find(qn(tag))
        if node is not None:
            pr.remove(node)
    pr.append(pr.makeelement(qn("a:noAutofit"), {}))


def _rgb(color: str) -> RGBColor:
    raw = (color or "#000000").strip().lstrip("#")
    if len(raw) != 6:
        raw = "000000"
    return RGBColor(int(raw[0:2], 16), int(raw[2:4], 16), int(raw[4:6], 16))
