"""Skill writer produces a real pptx with an editable OOXML chart when given series."""

from __future__ import annotations

import zipfile
from pathlib import Path

from pptx import Presentation

from app.generation.skill_writer import write_skill_pptx
from app.models.claim import ClaimLedger
from app.models.slide import ChartSeries, SlideChart, SlidePlan, SlidePlanSlide

FIXTURES = Path(__file__).resolve().parent / "fixtures"


def test_skill_writer_embeds_native_chart(tmp_path: Path) -> None:
    ledger = ClaimLedger.model_validate_json((FIXTURES / "ledger.json").read_text(encoding="utf-8"))
    plan = SlidePlan(
        blueprint_id="msl_physician_8",
        skill_version="v1",
        slides=[
            SlidePlanSlide(
                slide=1,
                role="primary_endpoint",
                headline="Primary endpoint — arms × outcomes",
                claim_ids=["C-014"],
                speaker_notes="C-014 p.4",
                chart=SlideChart(
                    title="Primary analysis",
                    categories=["Arm A", "Arm B", "Arm C"],
                    series=[
                        ChartSeries(name="Placebo", values=[12.0, 18.0, 9.0]),
                        ChartSeries(name="Active", values=[22.0, 31.0, 19.0]),
                    ],
                ),
            )
        ],
    )
    dest = write_skill_pptx(plan, ledger, tmp_path / "deck.pptx")
    assert dest.is_file()
    with zipfile.ZipFile(dest) as zf:
        names = zf.namelist()
    assert any(name.startswith("ppt/charts/") for name in names)
    assert any(name.startswith("ppt/embeddings/") for name in names)
    prs = Presentation(str(dest))
    assert len(prs.slides) == 1
    notes = prs.slides[0].notes_slide.notes_text_frame.text
    assert "C-014" in notes


def test_planner_none_args_still_stub() -> None:
    from app.generation import planner
    import pytest

    with pytest.raises(NotImplementedError):
        planner.plan_slides(blueprint=None, ledger=None, brief=None, skill_version="v1")
