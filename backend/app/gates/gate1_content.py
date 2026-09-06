"""Gate 1 — content & compliance checks (PRD §7.6 / Prompt 4)."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable

import yaml
from pptx import Presentation

from app.models.blueprint import Blueprint
from app.models.claim import ClaimLedger, ClaimType, EvidenceClass
from app.models.common import CheckLocation
from app.models.gates import Gate1Result, GateCheckResult
from app.models.numbers import NumbersIndex
from app.models.slide import SlideMap, SlideMapEntry
from app.paths import BLUEPRINTS_DIR, CONFIG_DIR

_CLAIM_ID_RE = re.compile(r"\bC-\d+\b", re.I)
_NUM_RE = re.compile(
    r"(?<![A-Za-z_])(?P<num>\d{1,3}(?:,\d{3})+(?:\.\d+)?|\d+\.\d+|\d+)(?P<unit>\s*%|\s*mg(?:/d[Ll])?|\s*g\b|\s*mmHg)?"
)
_LABELED_CLASSES = {
    EvidenceClass.SECONDARY_ENDPOINT,
    EvidenceClass.EXPLORATORY,
    EvidenceClass.POST_HOC,
}
_H2H_RE = re.compile(
    r"head[\s-]*to[\s-]*head|active\s+comparator|\bversus\b|\bvs\.?\b",
    re.I,
)
_FAIR_ROLES = ("safety", "limitations_relevance")


def check_claim_mapping(slide_map: SlideMap) -> GateCheckResult:
    """Every text run in ``slide_map`` must have ≥ 1 claim_id."""
    if slide_map is None:
        raise NotImplementedError
    locations: list[CheckLocation] = []
    for entry in slide_map.entries:
        if not entry.text.strip():
            continue
        if not entry.claim_ids:
            locations.extend(
                locations_for_element(entry.slide, entry.element_id, "no claim_id on text run")
            )
    passed = not locations
    return GateCheckResult(
        name="claim_mapping",
        passed=passed,
        message="All text runs mapped to claim IDs" if passed else f"{len(locations)} text run(s) missing claim_id",
        locations=locations,
    )


def check_number_sweep(
    pptx_path: Path | str,
    numbers_index: NumbersIndex,
    *,
    relative_tolerance: float,
    absolute_tolerance: float,
) -> GateCheckResult:
    """Every numeral in the deck must exist in ``numbers_index`` (with tolerances).

    Derived numbers are flagged (``flagged=True``), not failed.
    """
    allowed = _allowed_numbers(numbers_index)
    locations: list[CheckLocation] = []
    derived_locations: list[CheckLocation] = []
    prs = Presentation(str(pptx_path))
    for slide_no, element_id, text in _iter_visible_text(prs):
        for raw, value, unit in _extract_numerals(text):
            if _matches_index(value, unit, raw, allowed, relative_tolerance, absolute_tolerance):
                continue
            loc = locations_for_element(slide_no, element_id, f"numeral {raw}")
            if _is_derived(value, unit, raw, allowed, relative_tolerance, absolute_tolerance):
                derived_locations.extend(loc)
            else:
                locations.extend(loc)

    flagged = bool(derived_locations)
    passed = not locations
    if not passed:
        message = f"{len(locations)} numeral(s) not in numbers_index"
    elif flagged:
        message = f"{len(derived_locations)} derived numeral(s) flagged for human review"
    else:
        message = "All numerals found in numbers_index"
    return GateCheckResult(
        name="number_sweep",
        passed=passed,
        message=message,
        flagged=flagged,
        locations=locations + derived_locations,
    )


def check_endpoint_labelling(slide_map: SlideMap, ledger: ClaimLedger) -> GateCheckResult:
    """Secondary / exploratory / post-hoc claims must be labelled in slide text."""
    claims = {e.id: e for e in ledger.entries}
    labels = _compliance().get("endpoint_labels") or {
        "secondary_endpoint": "Secondary endpoint",
        "exploratory": "Exploratory",
        "post_hoc": "Post hoc",
    }
    by_slide = _entries_by_slide(slide_map)
    locations: list[CheckLocation] = []
    for slide_no, entries in by_slide.items():
        slide_text = " ".join(e.text for e in entries)
        slide_lower = slide_text.lower()
        seen: set[str] = set()
        for entry in entries:
            for cid in entry.claim_ids:
                if cid in seen:
                    continue
                seen.add(cid)
                claim = claims.get(cid)
                if claim is None or claim.evidence_class not in _LABELED_CLASSES:
                    continue
                label = str(labels.get(claim.evidence_class.value, claim.evidence_class.value))
                if label.lower() not in slide_lower:
                    locations.extend(
                        locations_for_element(
                            slide_no,
                            entry.element_id,
                            f"{cid} missing label {label!r}",
                        )
                    )
    passed = not locations
    return GateCheckResult(
        name="endpoint_labelling",
        passed=passed,
        message=(
            "No secondary/exploratory/post-hoc claims on deck"
            if passed and not any(
                (claims.get(cid) and claims[cid].evidence_class in _LABELED_CLASSES)
                for e in slide_map.entries
                for cid in e.claim_ids
            )
            else ("All non-primary endpoints labelled" if passed else f"{len(locations)} unlabelled endpoint claim(s)")
        ),
        locations=locations,
    )


def check_fair_balance(slide_map: SlideMap, blueprint_roles: list[str]) -> GateCheckResult:
    """Safety and limitations roles must be present and non-empty."""
    roles = list(blueprint_roles or [])
    required = list(_compliance().get("required_fair_balance_roles") or _FAIR_ROLES)
    by_slide = _entries_by_slide(slide_map)
    locations: list[CheckLocation] = []
    for role in required:
        idxs = [i + 1 for i, r in enumerate(roles) if r == role or (role == "limitations_relevance" and r.startswith("limitation"))]
        if not idxs:
            locations.append(CheckLocation(slide=1, detail=f"missing role {role}"))
            continue
        for slide_no in idxs:
            texts = [e.text.strip() for e in by_slide.get(slide_no, []) if e.text.strip()]
            if not texts:
                locations.append(
                    CheckLocation(slide=slide_no, detail=f"{role} slide is empty")
                )
    passed = not locations
    return GateCheckResult(
        name="fair_balance",
        passed=passed,
        message="Safety and limitations roles present" if passed else "Fair-balance role missing or empty",
        locations=locations,
    )


def check_comparative_claims(slide_map: SlideMap, ledger: ClaimLedger) -> GateCheckResult:
    """Comparative language requires a supporting head-to-head primary_endpoint claim."""
    phrases = list(_compliance().get("comparative_language") or [])
    patterns = [re.compile(re.escape(p), re.I) for p in phrases if p]
    has_h2h = _ledger_has_h2h_primary(ledger)
    locations: list[CheckLocation] = []
    for entry in slide_map.entries:
        for pat in patterns:
            if pat.search(entry.text):
                if not has_h2h:
                    locations.extend(
                        locations_for_element(
                            entry.slide,
                            entry.element_id,
                            f"comparative language without head-to-head primary_endpoint: {pat.pattern}",
                        )
                    )
                break
    passed = not locations
    return GateCheckResult(
        name="comparative_claims",
        passed=passed,
        message=(
            "No unsupported comparative language"
            if passed
            else f"{len(locations)} comparative phrase(s) without head-to-head primary_endpoint"
        ),
        locations=locations,
    )


def check_promotional_language(
    slide_map: SlideMap,
    blocklist: list[str],
) -> GateCheckResult:
    """Fail if any promotional blocklist term appears in deck text."""
    terms = [t for t in blocklist if t]
    locations: list[CheckLocation] = []
    for entry in slide_map.entries:
        lower = entry.text.lower()
        for term in terms:
            if term.lower() in lower:
                locations.extend(
                    locations_for_element(entry.slide, entry.element_id, f"blocklist term {term!r}")
                )
    passed = not locations
    return GateCheckResult(
        name="promotional_language",
        passed=passed,
        message="No blocklist terms" if passed else f"{len(locations)} promotional term(s) found",
        locations=locations,
    )


def check_population_extrapolation(slide_map: SlideMap, ledger: ClaimLedger) -> GateCheckResult:
    """No claim may generalize beyond the ledger's stated population."""
    patterns = [
        re.compile(re.escape(p), re.I)
        for p in (_compliance().get("population_extrapolation_patterns") or [])
        if p
    ]
    ledger_text = " ".join(f"{e.text} {e.verbatim}" for e in ledger.entries).lower()
    locations: list[CheckLocation] = []
    for entry in slide_map.entries:
        for pat in patterns:
            match = pat.search(entry.text)
            if not match:
                continue
            phrase = match.group(0)
            if phrase.lower() not in ledger_text:
                locations.extend(
                    locations_for_element(
                        entry.slide,
                        entry.element_id,
                        f"population extrapolation {phrase!r}",
                    )
                )
    passed = not locations
    return GateCheckResult(
        name="population_extrapolation",
        passed=passed,
        message="No unsupported population extrapolation" if passed else f"{len(locations)} extrapolation phrase(s)",
        locations=locations,
    )


def check_references_complete(slide_map: SlideMap, ledger: ClaimLedger) -> GateCheckResult:
    """Reference slide present; every cited claim listed."""
    del ledger
    by_slide = _entries_by_slide(slide_map)
    cited: set[str] = set()
    for entries in by_slide.values():
        for e in entries:
            cited.update(e.claim_ids)

    roles = _roles_for_map(slide_map)
    ref_slides = [i + 1 for i, role in enumerate(roles) if role == "references"]
    if not ref_slides:
        ref_slides = [
            slide
            for slide, entries in by_slide.items()
            if any("reference" in e.text.lower() for e in entries)
        ]
    if not ref_slides and cited:
        # Last slide that contains every cited id.
        ref_slides = [s for s, ents in by_slide.items() if cited <= {c for e in ents for c in e.claim_ids}]
        if ref_slides:
            ref_slides = [max(ref_slides)]

    locations: list[CheckLocation] = []
    if not ref_slides:
        locations.append(CheckLocation(slide=1, detail="reference slide missing"))
    else:
        listed: set[str] = set()
        for slide_no in ref_slides:
            for e in by_slide.get(slide_no, []):
                listed.update(e.claim_ids)
        missing = sorted(cited - listed)
        for cid in missing:
            locations.append(
                CheckLocation(slide=ref_slides[0], detail=f"cited claim {cid} not listed on references")
            )
        if not any(e.text.strip() for s in ref_slides for e in by_slide.get(s, [])):
            locations.append(CheckLocation(slide=ref_slides[0], detail="reference slide empty"))

    passed = not locations
    return GateCheckResult(
        name="references_complete",
        passed=passed,
        message="Reference slide lists every cited claim" if passed else "Reference slide incomplete",
        locations=locations,
    )


def check_synthesized_flags(slide_map: SlideMap, ledger: ClaimLedger) -> GateCheckResult:
    """Every ``synthesized`` claim is flagged for explicit human approval."""
    claims = {e.id: e for e in ledger.entries}
    locations: list[CheckLocation] = []
    seen: set[str] = set()
    for entry in slide_map.entries:
        for cid in entry.claim_ids:
            if cid in seen:
                continue
            claim = claims.get(cid)
            if claim is None or claim.claim_type != ClaimType.SYNTHESIZED:
                continue
            seen.add(cid)
            locations.extend(
                locations_for_element(entry.slide, entry.element_id, f"{cid} synthesized — needs human approval")
            )
    flagged = bool(locations)
    return GateCheckResult(
        name="synthesized_flags",
        passed=True,
        message="No synthesized claims" if not flagged else f"{len(locations)} synthesized claim(s) need human approval",
        flagged=flagged,
        locations=locations,
    )


def run_gate1(
    *,
    pptx_path: Path | str,
    slide_map: SlideMap,
    ledger: ClaimLedger,
    numbers_index: NumbersIndex,
    compliance_config: Path | str | None = None,
    blueprint_roles: list[str] | None = None,
) -> Gate1Result:
    """Run all Gate 1 checks and write-compatible ``gate1.json`` payload.

    Deterministic only — no LLM calls. Any failure blocks Gate 2.
    """
    cfg = _compliance(compliance_config)
    ns = cfg.get("number_sweep") or {}
    rel = float(ns.get("relative_tolerance", 0.01))
    abs_tol = float(ns.get("absolute_tolerance", 0.05))
    blocklist = list(cfg.get("promotional_blocklist") or [])
    roles = list(blueprint_roles) if blueprint_roles is not None else _roles_for_map(slide_map)

    checks = [
        check_claim_mapping(slide_map),
        check_number_sweep(
            pptx_path,
            numbers_index,
            relative_tolerance=rel,
            absolute_tolerance=abs_tol,
        ),
        check_endpoint_labelling(slide_map, ledger),
        check_fair_balance(slide_map, roles),
        check_comparative_claims(slide_map, ledger),
        check_promotional_language(slide_map, blocklist),
        check_population_extrapolation(slide_map, ledger),
        check_references_complete(slide_map, ledger),
        check_synthesized_flags(slide_map, ledger),
    ]
    return Gate1Result(gate="gate1", passed=all(c.passed for c in checks), checks=checks)


def locations_for_element(slide: int, element_id: str, detail: str = "") -> list[CheckLocation]:
    """Helper signature for building element-level failure locations."""
    return [CheckLocation(slide=slide, element_id=element_id, detail=detail)]


def _compliance(path: Path | str | None = None) -> dict[str, Any]:
    cfg_path = Path(path) if path is not None else CONFIG_DIR / "compliance.yaml"
    if not cfg_path.is_file():
        return {}
    data = yaml.safe_load(cfg_path.read_text(encoding="utf-8")) or {}
    return data if isinstance(data, dict) else {}


def _load_default_roles() -> list[str]:
    for name in ("msl_physician_8.yaml", "msl_physician_8.json"):
        path = BLUEPRINTS_DIR / name
        if not path.is_file():
            continue
        raw = path.read_text(encoding="utf-8")
        payload = yaml.safe_load(raw) if path.suffix in {".yaml", ".yml"} else __import__("json").loads(raw)
        blueprint = Blueprint.model_validate(payload)
        return [s.role for s in blueprint.slides]
    return list(_FAIR_ROLES) + ["references"]


def _roles_for_map(slide_map: SlideMap) -> list[str]:
    roles = _load_default_roles()
    max_slide = max((e.slide for e in slide_map.entries), default=len(roles))
    if len(roles) < max_slide:
        roles = roles + [""] * (max_slide - len(roles))
    return roles


def _entries_by_slide(slide_map: SlideMap) -> dict[int, list[SlideMapEntry]]:
    out: dict[int, list[SlideMapEntry]] = {}
    for entry in slide_map.entries:
        out.setdefault(entry.slide, []).append(entry)
    return out


def _ledger_has_h2h_primary(ledger: ClaimLedger) -> bool:
    primaries = [e for e in ledger.entries if e.evidence_class == EvidenceClass.PRIMARY_ENDPOINT]
    if any(_H2H_RE.search(f"{e.text} {e.verbatim}") for e in primaries):
        return True
    designs = [e for e in ledger.entries if e.evidence_class == EvidenceClass.DESIGN]
    return bool(primaries) and any(_H2H_RE.search(f"{e.text} {e.verbatim}") for e in designs)


def _iter_visible_text(prs) -> Iterable[tuple[int, str, str]]:
    for slide_no, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            name = shape.name or ""
            if getattr(shape, "has_text_frame", False):
                text = shape.text_frame.text
                if text.strip():
                    yield slide_no, name, text
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            yield slide_no, name, cell.text


def _extract_numerals(text: str) -> list[tuple[str, float, str | None]]:
    masked = _CLAIM_ID_RE.sub(" ", text)
    found: list[tuple[str, float, str | None]] = []
    for match in _NUM_RE.finditer(masked):
        raw_num = match.group("num")
        unit_raw = match.group("unit")
        unit = unit_raw.strip() if unit_raw else None
        try:
            value = float(raw_num.replace(",", ""))
        except ValueError:
            continue
        raw = (raw_num + (unit_raw or "")).strip()
        found.append((raw, value, unit))
    return found


def _allowed_numbers(index: NumbersIndex) -> list[tuple[float, str | None, str]]:
    allowed: list[tuple[float, str | None, str]] = []
    for entry in index.numbers:
        allowed.append((entry.value, entry.unit, entry.raw))
        for raw, value, unit in _extract_numerals(entry.raw):
            allowed.append((value, unit or entry.unit, raw))
        for raw, value, unit in _extract_numerals(entry.sentence):
            allowed.append((value, unit, raw))
    return allowed


def _variants(value: float, unit: str | None, raw: str) -> list[tuple[float, str | None]]:
    out = [(value, unit)]
    u = (unit or "").lower()
    raw_l = (raw or "").lower()
    if u in {"%", "percent", "pct"} or raw_l.rstrip().endswith("%"):
        out.append((value / 100.0, "fraction"))
    elif u in {"fraction", "prop", "proportion"}:
        out.append((value * 100.0, "%"))
    if u == "mg":
        out.append((value / 1000.0, "g"))
    elif u == "g":
        out.append((value * 1000.0, "mg"))
    return out


def _close(a: float, b: float, rel: float, abs_tol: float) -> bool:
    return abs(a - b) <= max(abs_tol, rel * max(abs(a), abs(b), 1e-12))


def _matches_index(
    value: float,
    unit: str | None,
    raw: str,
    allowed: list[tuple[float, str | None, str]],
    rel: float,
    abs_tol: float,
) -> bool:
    deck_vars = _variants(value, unit, raw)
    for aval, aunit, araw in allowed:
        for idx_val, _idx_unit in _variants(aval, aunit, araw):
            for dval, _dunit in deck_vars:
                if _close(dval, idx_val, rel, abs_tol):
                    return True
    return False


def _is_derived(
    value: float,
    unit: str | None,
    raw: str,
    allowed: list[tuple[float, str | None, str]],
    rel: float,
    abs_tol: float,
) -> bool:
    bases = [aval for aval, _u, _r in allowed]
    unique: list[float] = []
    for b in bases:
        if not any(_close(b, u, rel, abs_tol) for u in unique):
            unique.append(b)
    unique = unique[:12]
    deck_vars = [v for v, _u in _variants(value, unit, raw)]
    candidates: list[float] = []
    for a in unique:
        candidates.append(100.0 - a)
        for b in unique:
            candidates.extend((a + b, a - b, b - a, a * b))
            if abs(b) > 1e-12:
                candidates.append(a / b)
            if abs(a) > 1e-12:
                candidates.append(b / a)
    for dval in deck_vars:
        if any(_close(dval, c, rel, abs_tol) for c in candidates):
            return True
    return False
